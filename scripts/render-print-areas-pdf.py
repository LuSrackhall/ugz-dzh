#!/usr/bin/env python3
"""渲染打印版 ML 账页的打印区域为 PDF（修复前 vs 修复后 对比）。

用途：v0.10.1 打印分页修复的可视化验证——按 _xlnm.Print_Area 的区域列表
逐区域渲染成一页 PDF，忠实呈现"导出 PDF 时会打出哪些页、每页什么内容"。

- 「修复前」用旧 mlAreaPlan 算法（滑动窗口：末块只取左半）重算区域列表
- 「修复后」读打印版 xlsx 里实际写入的区域列表
两份 PDF 逐页对比，即可看出修复前缺失的"末页正面"。

用法：
  python3 scripts/render-print-areas-pdf.py <打印版xlsx> <sheet名> <输出目录>
"""
import re
import sys
import zipfile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries, get_column_letter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

# 中文字体（macOS 系统字体）
FONT = "STHeiti"
try:
    pdfmetrics.registerFont(TTFont(FONT, "/System/Library/Fonts/STHeiti Light.ttc"))
except Exception:
    try:
        pdfmetrics.registerFont(TTFont(FONT, "/System/Library/Fonts/PingFang.ttc"))
    except Exception:
        FONT = "Helvetica"  # 兜底（中文会缺字）


def parse_area_text(text):
    out = []
    for part in text.split(","):
        m = re.search(r"\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+)", part)
        c1, r1, c2, r2 = range_boundaries(f"{m.group(1)}{m.group(2)}:{m.group(3)}{m.group(4)}")
        out.append((c1, r1, c2, r2))
    return out


def read_area_xml(xlsx, sheet):
    zin = zipfile.ZipFile(xlsx)
    root = ET.fromstring(zin.read("xl/workbook.xml").decode("utf-8"))
    wb = load_workbook(xlsx, read_only=False)
    sid = str(wb.sheetnames.index(sheet))
    for dn in root.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}definedName"):
        if dn.get("name") == "_xlnm.Print_Area" and dn.get("localSheetId") == sid:
            return parse_area_text(dn.text)
    return []


def old_ml_area_plan(last_row, block_rows, break_col, max_col):
    """旧算法（v0.10.0 及以前，供对比）：滑动窗口，末块只取左半。"""
    blocks = (last_row + block_rows - 1) // block_rows

    def rect(left, k):
        r1 = k * block_rows + 1
        r2 = min(r1 + block_rows - 1, last_row)
        if left:
            return (1, r1, break_col - 1, r2)
        return (break_col, r1, max_col, r2)

    rects = [rect(False, 0)]
    for k in range(1, blocks - 1):
        rects.append(rect(True, k))
        rects.append(rect(False, k))
    if blocks >= 2:
        rects.append(rect(True, blocks - 1))
    return rects


def render(ws, regions, title, out_path, highlight=None, missing_note=None):
    """逐区域渲染：每个区域一页 PDF，按行列出区域内非空单元格。

    highlight: 需要标注"★ 本次修复新增"的区域坐标集合（区域元组）。
    missing_note: 末尾追加一页说明（如旧版缺失的区域列表）。
    """
    c = canvas.Canvas(out_path, pagesize=landscape(A4))
    W, H = landscape(A4)
    for idx, region in enumerate(regions, 1):
        c1, r1, c2, r2 = region
        rows_map = {}
        for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
            for cell in row:
                if cell.value is None:
                    continue
                v = str(cell.value).strip()
                if v == "":
                    continue
                rows_map.setdefault(cell.row, []).append((cell.column, v))
        tag = ""
        if highlight and region in highlight:
            tag = "   ★★★ 本次修复新增的页（修复前导出 PDF 缺失此页）★★★"
        header = f"{title} | 第 {idx}/{len(regions)} 页 | 区域 {get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}{tag}"
        c.setFont(FONT, 11)
        c.setFillColorRGB(0.8, 0, 0) if tag else c.setFillColorRGB(0, 0, 0)
        c.drawString(24, H - 26, header)
        c.setFillColorRGB(0, 0, 0)
        if not rows_map:
            c.setFont(FONT, 13)
            c.drawString(24, H / 2, "（空白补页 —— 保偶数配对用）")
            c.showPage()
            continue
        y = H - 46
        for r in sorted(rows_map):
            items = rows_map[r]
            shown = items[:18]
            seg = "  ".join(f"{get_column_letter(cc)}={v[:16]}" for cc, v in shown)
            if len(items) > len(shown):
                seg += f"  …(还有{len(items)-len(shown)}列)"
            c.setFont(FONT, 7)
            c.drawString(24, y, f"行{r:>3}: {seg}")
            y -= 12
            if y < 24:
                c.drawString(24, 12, f"...（该区域共 {len(rows_map)} 个内容行，已截断显示）")
                break
        c.showPage()
    if missing_note:
        c.setFont(FONT, 12)
        c.drawString(24, H - 40, missing_note[0])
        c.setFont(FONT, 9)
        y = H - 64
        for line in missing_note[1]:
            c.drawString(24, y, line)
            y -= 14
        c.showPage()
    c.save()


def main():
    xlsx, sheet, outdir = sys.argv[1], sys.argv[2], sys.argv[3]
    wb = load_workbook(xlsx, read_only=False)
    ws = wb[sheet]

    # 块参数（ML：30 行/块；分界列由区域列表推断）
    block_rows = 30
    new_regions = read_area_xml(xlsx, sheet)
    break_col = min(r[0] for r in new_regions if r[0] > 1)
    max_col = max(r[2] for r in new_regions)
    # 从新区域数反推生成时的块数（新算法区域数 = 2×blocks 含补空白；避免用
    # openpyxl max_row——它含尾部样式行，与生成期 GetRows 的 lastRow 不同）
    blocks = max(len(new_regions) // 2, 1)
    last_row = blocks * block_rows
    old_regions = old_ml_area_plan(last_row, block_rows, break_col, max_col)

    import os
    p_old = os.path.join(outdir, "打印版-修复前.pdf")
    p_new = os.path.join(outdir, "打印版-修复后.pdf")
    missing = [r for r in new_regions if r not in old_regions]
    note = (
        "⚠ 旧算法（修复前）缺失以下打印区域 —— 导出 PDF 时这些页不会出现：",
        [f"  * {get_column_letter(r[0])}{r[1]}:{get_column_letter(r[2])}{r[3]}" +
         ("（末页正面：标题/页码/会计科目/表头/明细5-14 数据）" if r[1] > 1 and r[0] > 1 else "")
         for r in missing],
    ) if missing else None
    render(ws, old_regions, f"{sheet}（修复前：旧区域算法）", p_old, missing_note=note)
    render(ws, new_regions, f"{sheet}（修复后：当前实现）", p_new, highlight=set(missing))
    print(f"修复前 {len(old_regions)} 页 → {p_old}")
    print(f"修复后 {len(new_regions)} 页 → {p_new}")
    print(f"块数 { -(-last_row // block_rows)}，行数 {last_row}，分界列 {get_column_letter(break_col)}({break_col})")


if __name__ == "__main__":
    main()
